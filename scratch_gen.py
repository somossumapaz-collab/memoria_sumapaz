import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>'))

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_after = Pt(2)
r_title = p_title.add_run('LOCALIDAD DE SUMAPAZ: CARACTERIZACIÓN SOCIOECONÓMICA Y PRODUCTIVA')
r_title.font.name = 'Calibri'
r_title.font.size = Pt(17)
r_title.font.bold = True
R_title.font.color.rgb = RGBColor(27, 94, 32)

p_sub = doc.add_paragraph()
p_sub.alignment = WD_TABLE_ALIGNMENT.CENTER
p_sub.paragraph_format.space_after = Pt(16)
r_sub = p_sub.add_run('Informe Técnico de Respuestas Basadas en la Base de Datos PMAPC (71 Unidades Productivas Caracterizadas)')
R_sub.font.name = 'Calibri'
r_sub.font.size = Pt(11)
r_sub.font.italic = True
R_sub.font.color.rgb = RGBColor(90, 90, 90)

h1 = doc.add_heading(level=1)
h1.paragraph_format.space_before = Pt(10)
h1.paragraph_format.space_after = Pt(8)
r_h1 = h1.add_run('PREGUNTA 1: Principales Actividades Económicas Desarrolladas en Sumapaz')
r_h1.font.name = 'Calibri'
r_h1.font.size = Pt(14)
r_h1.font.bold = True
R_h1.font.color.rgb = RGBColor(27, 94, 32)

tbl_q = doc.add_table(rows=1, cols=1)
tblq.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl_q.autofit = False
cell_q = tblq.cell(0, 0)
cellq.width = Inches(6.5)
set_cell_background(cell_q, 'F0F4F8')
set_cell_margins(cell_q, top=140, bottom=140, left=180, right=180)
p_q = cellq.paragraphs[0]
p_q.paragraph_format.space_after = Pt(0)
r_q_label = p_q.add_run('Pregunta formulada: ')
r_q_label.font.name = 'Calibri'
r_q_label.font.bold = True
r_q_label.font.size = Pt(10.5)
r_q_label.font.color.rgb = RGBColor(30, 60, 90)
r_q_text = p_q.add_run('“¿Cuáles son las principales actividades económicas desarrolladas actualmente en Sumapaz? Solicito información, si existe, sobre agricultura, ganaderÁa, producción lechera, transformación de alimentos, turismo rural, comercio y servicios.”')
R_q_text.font.name = 'Calibri'
r_q_text.font.italic = True
r_q_text.font.size = Pt(10.5)

pspace = doc.add_paragraph()
p_space.paragraph_format.space_after = Pt(6)

h2_1 = doc.add_heading(level=2)
h2_1.paragraph_format.space_before = Pt(10)
h2_1.paragraph_format.space_after = Pt(6)
r_h2_1 = h2_1.add_run('1. Resumen Ejecutivo (Respuesta Corta y Consistente)')
r_h2_1.font.name = 'Calibri'
r_h2_1.font.size = Pt(12)	�_h2_1.font.bold = True
R_h2_1.font.color.rgb = RGBColor(46, 125, 50)

p_intro = doc.add_paragraph()
p_intro.paragraph_format.line_spacing = 1.15
p_intro.paragraph_format.space_after = Pt(8)
r_intro = p_intro.add_run('A partir de la consolidación de la base de datos de 71 Unidades Productivas Caracterizadas (PMAPC) en la Localidad 20 de Sumapaz, la estructura económica territorial se caracteriza por un modelo dinámico de pluriactividad campesina. Las familias campesinas diversifican sus actividades productivas integrando la ganadería bovina de leche y ceba, la transformación agroindustrial artesanal en finca, la producción de especies menores, la manufactura textil tradicional y el comercio de servicios locales.')

bullets_summary = [
    ('Ganaderéa Bovina y Productividad Lechera (54.9% de las U.P. / 39 U.P.): ', 'Constituye el eje principal de liquidez econömica en la región. Predomina el ordeño diario con venta de leche cruda fresca en cantinas y acopios locales, articulada con la cria, levante y ceba de novillos/terneros de doble propósito.'),
    ('Transformación Agroindustrial Artesanal (50.7% de las U.P. / 36 U.P.): ', 'Agregación de valor en el predio mediante elaboración de queso campesino fresco prensado, cuajada, mantequilla artesanal, rico (derivado del suero lácteo), postres por capas con frutos silvestres, mantecadas campesinas y pomadas medicinales de caléndula.'),
    ('Especies Menores y Complementarias (39.4% de las U.P. / 28 U.P.): ', 'Sistemas de porcicultura (créa y levante de cerdos, entrega de lechones vivos a los 40 días i integración de biodigestores para biogás), avicultura (huevos campesinos en cubeta y pollo de engorde), ovinos y piscicultura (trucha arcoíris fresca por libra).'),
    ('Artesanéas, Textiles y Saberes Tradicionales (39.4% de las U.P. / 28 U.P.): ', 'Tejido artesanal en lana virgen de oveja y macramé (ruanas, gorros, bufandas, sacos, bolsos, amigurumis) con un fuerte componente de preservación de saberes y liderazgo de colectivos de mujeres (ej. veredas Tunal Alto y Tunal Bajo).'),
    ('Agricultura Agroecológica y de Páramo (23.9% de las U.P. / 17 U.P.): ', 'Cultivos adaptados al ecosistema de alta montaña: papa nativa (variedades corneta, negra, criolla) y convencional, arveja verde en vaina, frutales (mora en establecimiento, uchuva, tomate de árbol, arándanos), superalimentos (quinua y amaranto), tubérculos (cubios e ibias) y huertas agroecológicas de aromáticas.'),
    (']Comercio Minorista y Servicios Locales (19-7% de las U.P. / 14 U.P.): ', 'Puntos de avituallamiento, tiendas de vÁveres, salsamentarias, cafeterias/comidas rápidas, venta de vestuario térmico (sÁbanas/pijamas térmicas), insumos agrÃcolas y vacunas/medicamentos veterinarios, y servicios de mantenimiento e infraestructura.'),
    ('Turismo Rural y Alojamiento Campesino (1.4% de las U.P. PMAPC): ', 'Servicios de hospedaje y alojamiento rural por noches, semanas o meses en habitaciones acondicionadas en predios campesinos, integrados a recorridos agroturísticos por el páramo y oferta gastronó=ica local.')
]
